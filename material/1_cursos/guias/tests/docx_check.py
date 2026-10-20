"""Valida los .docx de tests/salida: XML bien formado y lectura con python-docx."""
import glob, os, sys, zipfile, xml.dom.minidom
import docx
bad = 0
for f in sorted(glob.glob(os.path.join(os.path.dirname(__file__), 'salida', '*.docx'))):
    z = zipfile.ZipFile(f)
    for n in z.namelist():
        if n.endswith('.xml') or n.endswith('.rels'):
            try:
                xml.dom.minidom.parseString(z.read(n))
            except Exception as e:
                bad += 1
                print('XML', f, n, e)
    d = docx.Document(f)
    h1 = sum(1 for p in d.paragraphs if p.style.name == 'Heading 1')
    print(os.path.basename(f), 'parrafos', len(d.paragraphs), 'tablas', len(d.tables), 'imagenes', len(d.inline_shapes), 'H1', h1)
print('errores XML', bad)
sys.exit(1 if bad else 0)
