"""Genera plantillas/cuadro-de-cargas.xlsx desde la definición de lib/60-hoja.js.

Uso: python tools/xlsx.py
Requiere Node (para leer la definición) y openpyxl.
Las celdas con fórmula se escriben como fórmula; Excel y LibreOffice las calculan al abrir.
"""
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, 'plantillas', 'cuadro-de-cargas.xlsx')

datos = json.loads(subprocess.run(['node', os.path.join(RAIZ, 'tools', 'datos-hoja.js')], capture_output=True, check=True).stdout.decode('utf-8'))

AZUL = '1F4E9A'
borde = Border(*(Side(style='thin', color='B8BCC0'),) * 4)
ESTILOS = {
    'titulo': dict(font=Font(bold=True, size=14)),
    'nota': dict(font=Font(italic=True, color='5F6368', size=9)),
    'cab': dict(font=Font(bold=True, size=10), fill=PatternFill('solid', fgColor='E8E6E1'), alignment=Alignment(wrap_text=True, vertical='center', horizontal='center'), border=borde),
    'dato': dict(font=Font(color=AZUL)),
    'total': dict(font=Font(bold=True)),
    'ref': dict(font=Font(color='5F6368', size=9)),
}
FORMATOS = {'0': '0', '0.0': '0.0', '0.00': '0.00', '0.000': '0.000', '0%': '0%', '0.0%': '0.0%'}

wb = Workbook()
wb.remove(wb.active)
for h in datos['hojas']:
    ws = wb.create_sheet(h['nombre'])
    for col, ancho in h['anchos'].items():
        ws.column_dimensions[col].width = ancho
    for ref, c in h['celdas'].items():
        cel = ws[ref]
        if 'f' in c:
            cel.value = c['f']
        else:
            cel.value = c.get('v')
        if c.get('t') == 'txt':
            cel.value = str(cel.value)
            cel.number_format = '@'
        for k, v in ESTILOS.get(c.get('s'), {}).items():
            setattr(cel, k, v)
        if c.get('n'):
            cel.number_format = FORMATOS[c['n']]
        if ws.title == 'Cuadro' and cel.row >= datos['fila0'] and cel.row <= datos['filaTot']:
            cel.border = borde

cu = wb['Cuadro']
cu.freeze_panes = 'C5'
cu.row_dimensions[4].height = 44
f0, ft = datos['fila0'], datos['filaTot']
awg = '"14,12,10,8,6,4,3,2,1"'
dv = DataValidation(type='list', formula1=awg, allow_blank=False, showErrorMessage=True, errorTitle='Calibre', error='Escoja un calibre de la hoja Tablas.')
dv.add(f'L{f0}:M{ft - 1}')
cu.add_data_validation(dv)
dvf = DataValidation(type='list', formula1='"A,B,AB"', allow_blank=False)
dvf.add(f'D{f0}:D{ft - 1}')
cu.add_data_validation(dvf)
cu.page_setup.orientation = 'landscape'
cu.page_setup.paperSize = cu.PAPERSIZE_A3
cu.page_setup.fitToWidth = 1
cu.sheet_properties.pageSetUpPr.fitToPage = True
ac = wb['Acometida']
dva = DataValidation(type='list', formula1=awg)
dva.add('B4')
ac.add_data_validation(dva)
wb.properties.creator = 'Mg. Manuel Alejandro Obando · Departamento de Física, Universidad del Cauca'
wb.properties.title = 'Cuadro de cargas · Cableado con criterio'
wb.active = 1
os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
wb.save(SALIDA)
print('OK', SALIDA, sum(len(h['celdas']) for h in datos['hojas']), 'celdas')
